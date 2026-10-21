#!/usr/bin/env python3
"""
McKinsey & Company HTML-to-PPTX Presentation Compiler
Converts McKinsey / Slideworks-styled HTML slide decks (<section class="slide">)
into 100% native, editable 16:9 PowerPoint (.pptx) presentations.

Supports:
- Dark Navy Cover Slide (slide--cover) with Georgia serif display & blue line motifs
- Section Breaker Slide (slide--divider)
- Column Chart + 1/3 Grey Takeaways Split Slide (slide--chart-takeaways)
- 3x3 Prioritization Matrix Slide (slide--matrix) with RAG status badges
- Four Key Trends / Pillars Table Slide (slide--trends)
- Capability Maturity & Readiness Assessment Slide (slide--assessment)
- Native PowerPoint Speaker Notes from `data-notes` attributes
"""

import argparse
import os
import re
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# MCKINSEY & SLIDEWORKS 2025 DESIGN TOKENS
# ==============================================================================
COLOR_NAVY = RGBColor(0x05, 0x1C, 0x2C)          # #051C2C Primary Deep McKinsey Navy
COLOR_NAVY_ALT = RGBColor(0x06, 0x1F, 0x32)      # #061F32 Dark Slide / Table Header Navy
COLOR_ELECTRIC_BLUE = RGBColor(0x22, 0x22, 0xF6) # #2222F6 Vivid Electric Blue
COLOR_CHART_BLUE = RGBColor(0x1F, 0x6F, 0xD8)    # #1F6FD8 Forecast / Highlight Series Blue
COLOR_CYAN = RGBColor(0x22, 0xA3, 0xDF)          # #22A3DF Secondary Data Cyan
COLOR_SKY = RGBColor(0x5A, 0xCB, 0xF0)           # #5ACBF0 Sky Blue Accent
COLOR_MID_BLUE = RGBColor(0x18, 0x77, 0xA2)      # #1877A2 Corporate Tracker Blue
COLOR_ICE_BLUE = RGBColor(0xE7, 0xF0, 0xF9)      # #E7F0F9 Light Ice Blue Container Fill
COLOR_ICE_BORDER = RGBColor(0xA2, 0xD5, 0xEC)    # #A2D5EC Ice Blue Border / Subtitle Light

COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)         # #FFFFFF Pure White
COLOR_BG_GREY = RGBColor(0xF2, 0xF2, 0xF2)       # #F2F2F2 1/3 Split Grey / Alternating Row
COLOR_HEADER_GREY = RGBColor(0xD8, 0xD8, 0xD8)   # #D8D8D8 Parameter Row / Border Grey
COLOR_AXIS_GREY = RGBColor(0xA5, 0xA5, 0xA5)     # #A5A5A5 Axis / Muted Label Grey
COLOR_MUTED = RGBColor(0x7F, 0x7F, 0x7F)         # #7F7F7F Footnotes & Subtitles

# RAG Status Accents
COLOR_GREEN = RGBColor(0x00, 0xB0, 0x50)         # #00B050 On-Track / Wave 1 Priority
COLOR_AMBER = RGBColor(0xFA, 0xD6, 0x88)         # #FAD688 Medium / Wave 2 Scale
COLOR_AMBER_DARK = RGBColor(0xF8, 0xBC, 0x3B)    # #F8BC3B Amber Dot / Border
COLOR_CORAL = RGBColor(0xE5, 0x54, 0x6C)         # #E5546C Action Required / Heavy Capex

FONT_SERIF = "Georgia"  # PowerPoint-safe substitute for McKinsey Bower
FONT_SANS = "Arial"     # PowerPoint-safe substitute for McKinsey Sans


def clean_text(text: str) -> str:
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).strip()


def set_slide_background(slide, color: RGBColor):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, s_tag):
    """Renders the McKinsey Top-Left Tracker, Full-Sentence Action Title, Subtitle, and Navy Rule."""
    header_div = s_tag.find(class_="slide-header")
    if not header_div:
        return 0.6

    tracker_el = header_div.find(class_="slide-tracker")
    title_el = header_div.find(class_="slide-title")
    subtitle_el = header_div.find(class_="slide-subtitle")

    # Tracker Eyebrow
    if tracker_el:
        tb_tr = slide.shapes.add_textbox(Inches(0.65), Inches(0.32), Inches(12.0), Inches(0.25))
        tf_tr = tb_tr.text_frame
        tf_tr.word_wrap = True
        tf_tr.margin_left = tf_tr.margin_top = tf_tr.margin_right = tf_tr.margin_bottom = 0
        p = tf_tr.paragraphs[0]
        p.text = clean_text(tracker_el.get_text()).upper()
        p.font.name = FONT_SANS
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_MID_BLUE

    # Full-Sentence Action Title
    if title_el:
        tb_ti = slide.shapes.add_textbox(Inches(0.65), Inches(0.54), Inches(12.0), Inches(0.62))
        tf_ti = tb_ti.text_frame
        tf_ti.word_wrap = True
        tf_ti.margin_left = tf_ti.margin_top = tf_ti.margin_right = tf_ti.margin_bottom = 0
        p = tf_ti.paragraphs[0]
        p.text = clean_text(title_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

    # Subtitle
    if subtitle_el:
        tb_sub = slide.shapes.add_textbox(Inches(0.65), Inches(1.12), Inches(12.0), Inches(0.28))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p = tf_sub.paragraphs[0]
        p.text = clean_text(subtitle_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_MUTED

    # Horizontal Navy Divider Rule
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.42), Inches(12.033), Inches(0.02)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_NAVY
    rule.line.fill.background()

    return 1.56


def add_footer(slide, s_tag, slide_idx: int):
    """Renders McKinsey Source/Footnote attribution on bottom-left and McKinsey wordmark + page on bottom-right."""
    footer_div = s_tag.find(class_="slide-footer")
    source_text = "Source: McKinsey & Company analysis"
    if footer_div:
        src_el = footer_div.find(class_="footer-sources")
        if src_el:
            source_text = clean_text(src_el.get_text())

    # Top hairline border above footer
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(7.02), Inches(12.033), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_HEADER_GREY
    line.line.fill.background()

    # Left source & footnote
    tb_l = slide.shapes.add_textbox(Inches(0.65), Inches(7.08), Inches(9.5), Inches(0.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = source_text
    p_l.font.name = FONT_SANS
    p_l.font.size = Pt(7.5)
    p_l.font.color.rgb = COLOR_MUTED

    # Right McKinsey wordmark + slide number
    tb_r = slide.shapes.add_textbox(Inches(10.2), Inches(7.06), Inches(2.48), Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_right = tf_r.margin_top = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT

    run_brand = p_r.add_run()
    run_brand.text = "McKinsey & Company   "
    run_brand.font.name = FONT_SERIF
    run_brand.font.size = Pt(8.5)
    run_brand.font.bold = True
    run_brand.font.color.rgb = COLOR_NAVY

    run_pg = p_r.add_run()
    run_pg.text = str(slide_idx)
    run_pg.font.name = FONT_SANS
    run_pg.font.size = Pt(8.5)
    run_pg.font.bold = True
    run_pg.font.color.rgb = COLOR_NAVY


def render_cover_slide(slide, s_tag):
    """Renders Layout 1: Dark Navy Cover Slide (slide--cover)."""
    set_slide_background(slide, COLOR_NAVY)

    # Subtle architectural electric-blue & cyan accent bars on the right
    bar1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.85), Inches(0), Inches(0.16), Inches(7.5))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = COLOR_ELECTRIC_BLUE
    bar1.line.fill.background()

    bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(13.01), Inches(0), Inches(0.323), Inches(7.5))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = COLOR_CYAN
    bar2.line.fill.background()

    # Eyebrow
    eyebrow_el = s_tag.find(class_="cover-eyebrow")
    if eyebrow_el:
        tb_eye = slide.shapes.add_textbox(Inches(0.85), Inches(0.85), Inches(10.0), Inches(0.35))
        tf_eye = tb_eye.text_frame
        p = tf_eye.paragraphs[0]
        p.text = clean_text(eyebrow_el.get_text()).upper()
        p.font.name = FONT_SANS
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN

    # Main Cover Title (Georgia Serif)
    title_el = s_tag.find(class_="cover-title")
    if title_el:
        tb_t = slide.shapes.add_textbox(Inches(0.85), Inches(2.05), Inches(10.2), Inches(1.5))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p = tf_t.paragraphs[0]
        p.text = clean_text(title_el.get_text())
        p.font.name = FONT_SERIF
        p.font.size = Pt(34)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    # Electric Blue Accent Rule
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(3.65), Inches(1.1), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_ELECTRIC_BLUE
    rule.line.fill.background()

    # Subtitle
    sub_el = s_tag.find(class_="cover-subtitle")
    if sub_el:
        tb_s = slide.shapes.add_textbox(Inches(0.85), Inches(3.88), Inches(9.5), Inches(1.1))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p = tf_s.paragraphs[0]
        p.text = clean_text(sub_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(13.5)
        p.font.color.rgb = COLOR_ICE_BORDER

    # Bottom Divider Line
    f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(6.15), Inches(11.4), Inches(0.015))
    f_line.fill.solid()
    f_line.fill.fore_color.rgb = COLOR_MID_BLUE
    f_line.line.fill.background()

    # Footer Metadata
    meta_el = s_tag.find(class_="cover-meta")
    if meta_el:
        tb_m = slide.shapes.add_textbox(Inches(0.85), Inches(6.32), Inches(8.5), Inches(0.75))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        p = tf_m.paragraphs[0]
        p.text = clean_text(meta_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_HEADER_GREY

    # Stacked McKinsey Wordmark
    tb_b = slide.shapes.add_textbox(Inches(9.8), Inches(6.25), Inches(2.45), Inches(0.8))
    tf_b = tb_b.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.RIGHT
    p_b.text = "McKinsey &\nCompany"
    p_b.font.name = FONT_SERIF
    p_b.font.size = Pt(14)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_WHITE


def render_divider_slide(slide, s_tag):
    """Renders Layout 2: Section Breaker Slide (slide--divider)."""
    set_slide_background(slide, COLOR_NAVY_ALT)

    # Left Electric Blue vertical bar
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ELECTRIC_BLUE
    accent.line.fill.background()

    eye_el = s_tag.find(class_="divider-eyebrow")
    if eye_el:
        tb_e = slide.shapes.add_textbox(Inches(0.9), Inches(2.15), Inches(10.5), Inches(0.35))
        tf_e = tb_e.text_frame
        p = tf_e.paragraphs[0]
        p.text = clean_text(eye_el.get_text()).upper()
        p.font.name = FONT_SANS
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN

    title_el = s_tag.find(class_="divider-title")
    if title_el:
        tb_t = slide.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(10.8), Inches(1.6))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p = tf_t.paragraphs[0]
        p.text = clean_text(title_el.get_text())
        p.font.name = FONT_SERIF
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    sub_el = s_tag.find(class_="divider-sub")
    if sub_el:
        tb_s = slide.shapes.add_textbox(Inches(0.9), Inches(4.35), Inches(10.2), Inches(0.8))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p = tf_s.paragraphs[0]
        p.text = clean_text(sub_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(12.5)
        p.font.color.rgb = COLOR_ICE_BORDER

    # Footer
    tb_f = slide.shapes.add_textbox(Inches(0.9), Inches(6.75), Inches(11.5), Inches(0.4))
    tf_f = tb_f.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.alignment = PP_ALIGN.RIGHT
    p_f.text = "McKinsey & Company"
    p_f.font.name = FONT_SERIF
    p_f.font.size = Pt(10.5)
    p_f.font.bold = True
    p_f.font.color.rgb = COLOR_WHITE


def render_chart_takeaways_slide(slide, s_tag, top_y: float):
    """Renders Layout 3: Column Chart (Actuals vs Forecast) + 1/3 Grey Key Takeaways Sidebar."""
    # Left 2/3 Chart Panel
    chart_unit_el = s_tag.find(class_="chart-unit")
    chart_cagr_el = s_tag.find(class_="chart-cagr")

    if chart_unit_el:
        tb_u = slide.shapes.add_textbox(Inches(0.65), Inches(top_y), Inches(4.5), Inches(0.3))
        p = tb_u.text_frame.paragraphs[0]
        p.text = clean_text(chart_unit_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

    if chart_cagr_el:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.35), Inches(top_y - 0.02), Inches(2.35), Inches(0.32)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_ICE_BLUE
        badge.line.color.rgb = COLOR_ICE_BORDER
        p = badge.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = clean_text(chart_cagr_el.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

    # Native Drawn Column Bars from `.bar-group`
    bar_groups = s_tag.find_all(class_="bar-group")
    if bar_groups:
        max_val = max(float(bg.get("data-value", "100")) for bg in bar_groups)
        chart_bottom = 5.85
        max_bar_h = 3.25
        chart_left = 0.85
        chart_width = 6.65
        slot_w = chart_width / len(bar_groups)
        bar_w = min(0.95, slot_w * 0.56)

        # Baseline axis line
        axis_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(chart_bottom), Inches(7.0), Inches(0.02)
        )
        axis_line.fill.solid()
        axis_line.fill.fore_color.rgb = COLOR_NAVY
        axis_line.line.fill.background()

        for idx, bg in enumerate(bar_groups):
            val = float(bg.get("data-value", "0"))
            label = bg.get("data-label", f"Cat {idx+1}")
            sub = bg.get("data-sub", "")
            series = bg.get("data-series", "Forecast")
            delta = bg.get("data-delta", "")

            bar_h = max(0.25, (val / max_val) * max_bar_h)
            center_x = chart_left + idx * slot_w + (slot_w / 2.0)
            bar_x = center_x - (bar_w / 2.0)
            bar_y = chart_bottom - bar_h

            bar_shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(bar_x), Inches(bar_y), Inches(bar_w), Inches(bar_h)
            )
            bar_shape.fill.solid()
            if series == "Actuals":
                bar_shape.fill.fore_color.rgb = COLOR_NAVY
            elif series == "Target":
                bar_shape.fill.fore_color.rgb = COLOR_CYAN
            else:
                bar_shape.fill.fore_color.rgb = COLOR_CHART_BLUE
            bar_shape.line.fill.background()

            # Value label above bar
            tb_v = slide.shapes.add_textbox(
                Inches(center_x - 0.7), Inches(bar_y - 0.48), Inches(1.4), Inches(0.45)
            )
            tf_v = tb_v.text_frame
            tf_v.word_wrap = True
            p_d = tf_v.paragraphs[0]
            p_d.alignment = PP_ALIGN.CENTER
            p_d.text = delta
            p_d.font.name = FONT_SANS
            p_d.font.size = Pt(8)
            p_d.font.bold = True
            p_d.font.color.rgb = COLOR_CHART_BLUE

            p_v = tf_v.add_paragraph()
            p_v.alignment = PP_ALIGN.CENTER
            p_v.text = f"{int(val):,}" if val.is_integer() else f"{val}"
            p_v.font.name = FONT_SANS
            p_v.font.size = Pt(10.5)
            p_v.font.bold = True
            p_v.font.color.rgb = COLOR_NAVY

            # Category label below axis
            tb_lbl = slide.shapes.add_textbox(
                Inches(center_x - 0.8), Inches(chart_bottom + 0.06), Inches(1.6), Inches(0.55)
            )
            tf_lbl = tb_lbl.text_frame
            tf_lbl.word_wrap = True
            p_l1 = tf_lbl.paragraphs[0]
            p_l1.alignment = PP_ALIGN.CENTER
            p_l1.text = label
            p_l1.font.name = FONT_SANS
            p_l1.font.size = Pt(9.5)
            p_l1.font.bold = True
            p_l1.font.color.rgb = COLOR_NAVY

            if sub:
                p_l2 = tf_lbl.add_paragraph()
                p_l2.alignment = PP_ALIGN.CENTER
                p_l2.text = sub
                p_l2.font.name = FONT_SANS
                p_l2.font.size = Pt(8)
                p_l2.font.color.rgb = COLOR_MUTED

    # Legend at bottom of chart panel
    legend_items = [
        ("Actuals / Baseline", COLOR_NAVY),
        ("Wave 1–2 Forecast", COLOR_CHART_BLUE),
        ("Wave 3 Full Run-Rate Target", COLOR_CYAN),
    ]
    leg_x = 0.75
    for leg_text, leg_color in legend_items:
        sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(leg_x), Inches(6.55), Inches(0.14), Inches(0.14))
        sw.fill.solid()
        sw.fill.fore_color.rgb = leg_color
        sw.line.fill.background()

        tb_lg = slide.shapes.add_textbox(Inches(leg_x + 0.18), Inches(6.50), Inches(2.1), Inches(0.25))
        p_lg = tb_lg.text_frame.paragraphs[0]
        p_lg.text = leg_text
        p_lg.font.name = FONT_SANS
        p_lg.font.size = Pt(8.5)
        p_lg.font.color.rgb = COLOR_NAVY
        leg_x += 2.35

    # Right 1/3 Grey Key Takeaways Sidebar (Slideworks `1/3 grey` pattern)
    sidebar = s_tag.find(class_="takeaways-sidebar")
    if sidebar:
        sb_x = 8.05
        sb_w = 4.633
        sb_h = 5.15

        # Grey Background Box
        sb_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(sb_x), Inches(top_y), Inches(sb_w), Inches(sb_h)
        )
        sb_bg.fill.solid()
        sb_bg.fill.fore_color.rgb = COLOR_BG_GREY
        sb_bg.line.color.rgb = COLOR_HEADER_GREY

        # Dark Navy Header Bar
        sb_hdr = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(sb_x), Inches(top_y), Inches(sb_w), Inches(0.42)
        )
        sb_hdr.fill.solid()
        sb_hdr.fill.fore_color.rgb = COLOR_NAVY_ALT
        sb_hdr.line.fill.background()

        hdr_el = sidebar.find(class_="sidebar-header")
        p_sh = sb_hdr.text_frame.paragraphs[0]
        p_sh.text = clean_text(hdr_el.get_text()) if hdr_el else "Key takeaways / main conclusion"
        p_sh.font.name = FONT_SANS
        p_sh.font.size = Pt(11)
        p_sh.font.bold = True
        p_sh.font.color.rgb = COLOR_WHITE

        # Bullets inside sidebar
        tb_sb = slide.shapes.add_textbox(
            Inches(sb_x + 0.15), Inches(top_y + 0.52), Inches(sb_w - 0.3), Inches(sb_h - 0.6)
        )
        tf_sb = tb_sb.text_frame
        tf_sb.word_wrap = True
        for idx, li in enumerate(sidebar.find_all("li")):
            p = tf_sb.paragraphs[0] if idx == 0 else tf_sb.add_paragraph()
            p.space_after = Pt(10)
            bullet_run = p.add_run()
            bullet_run.text = "▪ "
            bullet_run.font.name = FONT_SANS
            bullet_run.font.size = Pt(10)
            bullet_run.font.bold = True
            bullet_run.font.color.rgb = COLOR_CHART_BLUE

            strong = li.find(["strong", "b"])
            if strong:
                lead_text = clean_text(strong.get_text())
                full_text = clean_text(li.get_text())
                rest_text = full_text[len(lead_text):].strip()

                r_bold = p.add_run()
                r_bold.text = lead_text + " "
                r_bold.font.name = FONT_SANS
                r_bold.font.size = Pt(10)
                r_bold.font.bold = True
                r_bold.font.color.rgb = COLOR_NAVY

                r_norm = p.add_run()
                r_norm.text = rest_text
                r_norm.font.name = FONT_SANS
                r_norm.font.size = Pt(10)
                r_norm.font.color.rgb = COLOR_NAVY
            else:
                r_norm = p.add_run()
                r_norm.text = clean_text(li.get_text())
                r_norm.font.name = FONT_SANS
                r_norm.font.size = Pt(10)
                r_norm.font.color.rgb = COLOR_NAVY


def render_matrix_slide(slide, s_tag, top_y: float):
    """Renders Layout 4: 3x3 Prioritization Matrix (Level of Impact vs Time to Impact) with RAG Badges."""
    grid_left = 0.65
    row_hdr_w = 1.35
    col_w = (12.033 - row_hdr_w - 0.18) / 3.0
    col_hdr_h = 0.36
    row_h = 1.38
    gap = 0.06

    # Corner label
    corner = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(grid_left), Inches(top_y), Inches(row_hdr_w), Inches(col_hdr_h)
    )
    corner.fill.solid()
    corner.fill.fore_color.rgb = COLOR_BG_GREY
    corner.line.color.rgb = COLOR_HEADER_GREY
    p_c = corner.text_frame.paragraphs[0]
    p_c.alignment = PP_ALIGN.CENTER
    p_c.text = "IMPACT vs. TIME"
    p_c.font.name = FONT_SANS
    p_c.font.size = Pt(8)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_MUTED

    col_headers = ["SHORT (< 6 Months)", "MEDIUM (6–15 Months)", "LONG (15–24+ Months)"]
    for c_idx, c_text in enumerate(col_headers):
        cx = grid_left + row_hdr_w + gap + c_idx * (col_w + gap)
        ch = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(cx), Inches(top_y), Inches(col_w), Inches(col_hdr_h)
        )
        ch.fill.solid()
        ch.fill.fore_color.rgb = COLOR_NAVY_ALT
        ch.line.fill.background()
        p = ch.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = c_text
        p.font.name = FONT_SANS
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    row_headers = [
        ("HIGH", "(> INR 100 Cr)"),
        ("MEDIUM", "(INR 40–100 Cr)"),
        ("LOW", "(< INR 40 Cr)"),
    ]
    cells = s_tag.find_all(class_="matrix-cell")

    for r_idx, (r_title, r_sub) in enumerate(row_headers):
        ry = top_y + col_hdr_h + gap + r_idx * (row_h + gap)
        rh = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(grid_left), Inches(ry), Inches(row_hdr_w), Inches(row_h)
        )
        rh.fill.solid()
        rh.fill.fore_color.rgb = COLOR_HEADER_GREY
        rh.line.fill.background()
        tf_r = rh.text_frame
        tf_r.word_wrap = True
        p1 = tf_r.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = r_title
        p1.font.name = FONT_SANS
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY

        p2 = tf_r.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = r_sub
        p2.font.name = FONT_SANS
        p2.font.size = Pt(8)
        p2.font.color.rgb = COLOR_NAVY

        for c_idx in range(3):
            cell_idx = r_idx * 3 + c_idx
            cx = grid_left + row_hdr_w + gap + c_idx * (col_w + gap)
            cell_el = cells[cell_idx] if cell_idx < len(cells) else None
            is_priority = cell_el and ("matrix-cell--priority" in cell_el.get("class", []))

            cell_box = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(cx), Inches(ry), Inches(col_w), Inches(row_h)
            )
            cell_box.fill.solid()
            cell_box.fill.fore_color.rgb = COLOR_ICE_BLUE if is_priority else COLOR_BG_GREY
            cell_box.line.color.rgb = COLOR_CHART_BLUE if is_priority else COLOR_HEADER_GREY

            if cell_el:
                badges = cell_el.find_all(class_="matrix-badge")
                for b_idx, b_el in enumerate(badges):
                    b_classes = b_el.get("class", [])
                    by = ry + 0.1 + b_idx * 0.58
                    b_shape = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE,
                        Inches(cx + 0.1),
                        Inches(by),
                        Inches(col_w - 0.2),
                        Inches(0.48),
                    )
                    b_shape.fill.solid()
                    if "badge--green" in b_classes:
                        b_shape.fill.fore_color.rgb = COLOR_GREEN
                        txt_color = COLOR_WHITE
                    elif "badge--coral" in b_classes:
                        b_shape.fill.fore_color.rgb = COLOR_CORAL
                        txt_color = COLOR_WHITE
                    else:
                        b_shape.fill.fore_color.rgb = COLOR_AMBER
                        txt_color = COLOR_NAVY
                    b_shape.line.fill.background()

                    spans = b_el.find_all("span")
                    if len(spans) >= 2:
                        badge_txt = f"{clean_text(spans[0].get_text())}  |  {clean_text(spans[1].get_text())}"
                    else:
                        badge_txt = clean_text(b_el.get_text())

                    p_b = b_shape.text_frame.paragraphs[0]
                    p_b.alignment = PP_ALIGN.CENTER
                    p_b.text = badge_txt
                    p_b.font.name = FONT_SANS
                    p_b.font.size = Pt(8.5)
                    p_b.font.bold = True
                    p_b.font.color.rgb = txt_color

    # Matrix Status Legend at bottom
    leg_y = top_y + col_hdr_h + 3 * (row_h + gap) + 0.12
    legends = [
        (COLOR_GREEN, "= Wave 1 Priority Quick Wins (Immediate Launch)"),
        (COLOR_AMBER, "= Wave 2 Core Scale Initiatives"),
        (COLOR_CORAL, "= Wave 3 Foundational / Heavy IT-OT Capex"),
    ]
    lx = 2.6
    for l_col, l_txt in legends:
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lx), Inches(leg_y), Inches(0.22), Inches(0.16))
        pill.fill.solid()
        pill.fill.fore_color.rgb = l_col
        pill.line.fill.background()

        tb_l = slide.shapes.add_textbox(Inches(lx + 0.26), Inches(leg_y - 0.05), Inches(3.2), Inches(0.28))
        p = tb_l.text_frame.paragraphs[0]
        p.text = l_txt
        p.font.name = FONT_SANS
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY
        lx += 3.35


def render_table_slide(slide, s_tag, top_y: float):
    """Renders Layout 5 (Four Key Trends Table) & Layout 6 (Capability Assessment Table)."""
    table_el = s_tag.find("table")
    if not table_el:
        return

    rows = table_el.find_all("tr")
    if not rows:
        return

    is_trends = "trends-table" in table_el.get("class", [])
    num_rows = len(rows)
    num_cols = max(len(r.find_all(["th", "td"])) for r in rows)

    table_h = 5.1 if is_trends else 4.35
    t_shape = slide.shapes.add_table(
        num_rows, num_cols, Inches(0.65), Inches(top_y), Inches(12.033), Inches(table_h)
    )
    table = t_shape.table

    # Column widths
    if is_trends and num_cols == 3:
        table.columns[0].width = Inches(2.75)
        table.columns[1].width = Inches(5.683)
        table.columns[2].width = Inches(3.60)
    elif num_cols == 4:
        table.columns[0].width = Inches(2.65)
        table.columns[1].width = Inches(4.583)
        table.columns[2].width = Inches(2.90)
        table.columns[3].width = Inches(1.90)

    for r_idx, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        is_header = (r_idx == 0) or bool(row.find("th"))
        for c_idx, cell_el in enumerate(cells):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()

            if is_header:
                cell.fill.fore_color.rgb = COLOR_NAVY_ALT
                cell.text = clean_text(cell_el.get_text())
                for p in cell.text_frame.paragraphs:
                    p.font.name = FONT_SANS
                    p.font.size = Pt(9.5)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_WHITE
            elif is_trends:
                if c_idx == 0:
                    cell.fill.fore_color.rgb = COLOR_NAVY
                    cell.text = clean_text(cell_el.get_text())
                    for p in cell.text_frame.paragraphs:
                        p.font.name = FONT_SANS
                        p.font.size = Pt(9.5)
                        p.font.bold = True
                        p.font.color.rgb = COLOR_WHITE
                else:
                    cell.fill.fore_color.rgb = COLOR_BG_GREY if c_idx == 1 else COLOR_ICE_BLUE
                    lis = cell_el.find_all("li")
                    tf = cell.text_frame
                    tf.word_wrap = True
                    if lis:
                        for li_idx, li in enumerate(lis):
                            p = tf.paragraphs[0] if li_idx == 0 else tf.add_paragraph()
                            p.text = f"• {clean_text(li.get_text())}"
                            p.font.name = FONT_SANS
                            p.font.size = Pt(8.8)
                            p.font.color.rgb = COLOR_NAVY
                    else:
                        cell.text = clean_text(cell_el.get_text())
            else:
                # Assessment table (Slideworks p96 alternating #D8D8D8 and #F2F2F2)
                if c_idx == 3:
                    txt = clean_text(cell_el.get_text())
                    if "High" in txt or "On Track" in txt:
                        cell.fill.fore_color.rgb = COLOR_GREEN
                        fg = COLOR_WHITE
                    elif "Low" in txt or "Action" in txt:
                        cell.fill.fore_color.rgb = COLOR_CORAL
                        fg = COLOR_WHITE
                    else:
                        cell.fill.fore_color.rgb = COLOR_AMBER
                        fg = COLOR_NAVY
                    cell.text = txt
                    for p in cell.text_frame.paragraphs:
                        p.alignment = PP_ALIGN.CENTER
                        p.font.name = FONT_SANS
                        p.font.size = Pt(9)
                        p.font.bold = True
                        p.font.color.rgb = fg
                else:
                    cell.fill.fore_color.rgb = COLOR_HEADER_GREY if (r_idx % 2 == 1) else COLOR_BG_GREY
                    cell.text = clean_text(cell_el.get_text())
                    for p in cell.text_frame.paragraphs:
                        p.font.name = FONT_SANS
                        p.font.size = Pt(9)
                        p.font.bold = (c_idx == 0)
                        p.font.color.rgb = COLOR_NAVY

    # Optional bottom callout bar (for Assessment slide)
    callout = s_tag.find(class_="bottom-callout-bar")
    if callout:
        c_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(6.22), Inches(12.033), Inches(0.62)
        )
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = COLOR_ICE_BLUE
        c_box.line.color.rgb = COLOR_ELECTRIC_BLUE
        p = c_box.text_frame.paragraphs[0]
        p.text = clean_text(callout.get_text())
        p.font.name = FONT_SANS
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY


def compile_html_to_pptx(html_path: str, pptx_path: str):
    if not os.path.exists(html_path):
        raise FileNotFoundError(f"Input HTML file not found: {html_path}")

    with open(html_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slides = soup.find_all("section", class_="slide")
    total_slides = len(slides)
    print(f"Found {total_slides} McKinsey slides in {html_path}")

    for idx, s_tag in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank_layout)
        classes = s_tag.get("class", [])

        if "slide--cover" in classes:
            render_cover_slide(slide, s_tag)
        elif "slide--divider" in classes:
            render_divider_slide(slide, s_tag)
        else:
            set_slide_background(slide, COLOR_WHITE)
            top_y = add_header(slide, s_tag)

            if "slide--chart-takeaways" in classes or s_tag.find(class_="split-2-1"):
                render_chart_takeaways_slide(slide, s_tag, top_y)
            elif "slide--matrix" in classes or s_tag.find(class_="matrix-grid"):
                render_matrix_slide(slide, s_tag, top_y)
            elif s_tag.find("table"):
                render_table_slide(slide, s_tag, top_y)

            add_footer(slide, s_tag, idx)

        # Native Speaker Notes
        notes_text = clean_text(s_tag.get("data-notes", ""))
        if notes_text:
            slide.notes_slide.notes_text_frame.text = notes_text

        print(f"  ✓ Compiled Slide {idx}/{total_slides}: {classes}")

    os.makedirs(os.path.dirname(os.path.abspath(pptx_path)), exist_ok=True)
    prs.save(pptx_path)
    print(f"\nSuccessfully generated McKinsey PowerPoint deck: {pptx_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Compile McKinsey HTML slide deck into editable PowerPoint (.pptx)")
    parser.add_argument("--input", "-i", required=True, help="Path to input McKinsey HTML file")
    parser.add_argument("--output", "-o", required=True, help="Path to output PowerPoint .pptx file")
    args = parser.parse_args()
    compile_html_to_pptx(args.input, args.output)
